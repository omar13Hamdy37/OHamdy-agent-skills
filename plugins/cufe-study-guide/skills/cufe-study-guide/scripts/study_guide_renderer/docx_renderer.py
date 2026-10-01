"""Word authoring with structural headings and explicit OOXML relationships."""

import hashlib
import io

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

from .cover_art import raster
from .equations import fitted
from .model import image_path, plain, runs
from .plan import navigation
from .theme import table_weights


def bookmark_name(identifier):
    return "b_" + hashlib.sha1(identifier.encode()).hexdigest()[:24]


def element(tag, **attrs):
    node = OxmlElement("w:" + tag)
    for name, value in attrs.items():
        node.set(qn("w:" + name), str(value))
    return node


def shade(paragraph, color, accent=None):
    properties = paragraph._p.get_or_add_pPr()
    properties.append(element("shd", fill=color.lstrip("#"), val="clear"))
    if accent:
        border = element("pBdr")
        border.append(element("left", val="single", sz=12, space=8, color=accent.lstrip("#")))
        properties.append(border)


class WordRenderer:
    def __init__(self, model, theme, equations, base):
        self.model, self.theme, self.equations, self.base = model, theme, equations, base
        self.document = Document()
        self.bookmark_id = 0
        styles = self.document.styles
        normal = styles["Normal"]
        normal.font.name = "DejaVu Sans"
        normal.font.size = Pt(theme["body_size"])
        normal.font.color.rgb = RGBColor.from_string(theme["text"][1:])
        normal.paragraph_format.line_spacing = 1.4
        normal.paragraph_format.space_after = Pt(7)
        normal.paragraph_format.widow_control = True
        for level in range(1, 7):
            style = styles[f"Heading {level}"]
            style.font.name = "DejaVu Sans"
            style.font.size = Pt((21, 16, 13, 12, 11, 11)[level - 1])
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.font.bold = True
            style.paragraph_format.space_before = Pt(18 if level == 1 else 12)
            style.paragraph_format.space_after = Pt(7)
            style.paragraph_format.keep_with_next = True
        for name, size in (("Title", 29), ("Subtitle", 13), ("Caption", 9)):
            style = styles[name]
            style.font.name = "DejaVu Sans"
            style.font.size = Pt(size)
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.paragraph_format.line_spacing = 1.2
        styles["Title"].font.bold = True
        for section in self.document.sections:
            self.geometry(section)
        self.document.core_properties.title = model["metadata"].get("title_override", model["metadata"]["title"])
        self.document.core_properties.subject = model["metadata"].get("course", "")
        self.document.core_properties.author = ""
        self.document.core_properties.last_modified_by = ""
        self.document.core_properties.created = self.document.core_properties.modified = __import__("datetime").datetime(2000, 1, 1)
        self.document.settings.element.append(element("updateFields", val="true"))
        lang = element("lang", val=model["metadata"].get("output_language", "en"))
        normal.element.get_or_add_rPr().append(lang)

    def geometry(self, section):
        section.page_width, section.page_height = map(Pt, self.theme["page"])
        section.top_margin = section.bottom_margin = Pt(self.theme["margin"])
        section.left_margin = section.right_margin = Pt(self.theme["margin"])
        section.header_distance = section.footer_distance = Pt(27)

    def anchor(self, paragraph, identifier):
        if not identifier:
            return
        self.bookmark_id += 1
        position = 1 if len(paragraph._p) and paragraph._p[0].tag == qn("w:pPr") else 0
        paragraph._p.insert(position, element("bookmarkStart", id=self.bookmark_id, name=bookmark_name(identifier)))
        paragraph._p.append(element("bookmarkEnd", id=self.bookmark_id))

    def rich(self, paragraph, value):
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        for item in runs(value):
            run = paragraph.add_run(item["text"])
            run.bold, run.italic = item.get("bold", False), item.get("italic", False)
            if item.get("code"):
                run.font.name = "DejaVu Sans Mono"
                run.font.size = Pt(10)
            if item.get("href"):
                href = item["href"]
                link = OxmlElement("w:hyperlink")
                if href.startswith("#"):
                    link.set(qn("w:anchor"), bookmark_name(href[1:]))
                else:
                    relation = paragraph.part.relate_to(href, RT.HYPERLINK, is_external=True)
                    link.set(qn("r:id"), relation)
                run.font.color.rgb = RGBColor.from_string(self.theme["link"][1:])
                run.underline = True
                link.append(run._r)
                paragraph._p.append(link)

    def cover(self):
        meta = self.model["metadata"]
        first = self.document.add_paragraph()
        first.paragraph_format.space_after = Pt(45)
        if meta.get("course"):
            p = self.document.add_paragraph(meta["course"], "Subtitle")
            p.paragraph_format.space_after = Pt(12)
        if "lecture_number" in meta:
            self.document.add_paragraph(f"Lecture {meta['lecture_number']}", "Subtitle")
        p = self.document.add_paragraph(meta.get("title_override", meta["title"]), "Title")
        p.paragraph_format.space_after = Pt(18)
        if len(p.text) > 110:
            for run in p.runs:
                run.font.size = Pt(24)
        if meta.get("subtitle"):
            self.document.add_paragraph(meta["subtitle"], "Subtitle")
        artwork = self.document.add_paragraph()
        artwork.paragraph_format.space_before = Pt(30)
        artwork.add_run().add_picture(io.BytesIO(raster(meta.get("subject_category", "general"), self.theme["link"])), width=Pt(self.theme["width"] * 0.85))
        for field in meta.get("cover_fields", []):
            self.document.add_paragraph(f"{field['label']}: {field['value']}")
        body = self.document.add_section(WD_SECTION_START.NEW_PAGE)
        self.geometry(body)
        body.header.is_linked_to_previous = body.footer.is_linked_to_previous = False
        header = body.header.paragraphs[0]
        header.text = meta.get("course", meta.get("title_override", meta["title"]))[:75]
        header.style = self.document.styles["Caption"]
        footer = body.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer.add_run()._r.append(element("fldChar", fldCharType="begin"))
        field = element("instrText")
        field.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        field.text = " PAGE "
        footer.add_run()._r.append(field)
        footer.add_run()._r.append(element("fldChar", fldCharType="separate"))
        footer.add_run("1")
        footer.add_run()._r.append(element("fldChar", fldCharType="end"))
        body._sectPr.append(element("pgNumType", start=1))

    def paragraph(self, value, semantic=None, style=None):
        p = self.document.add_paragraph(style=style)
        self.rich(p, value)
        if semantic and semantic not in ("quiz", "answer_key", "cheatsheet"):
            p.paragraph_format.left_indent = Pt(12)
        return p

    def table(self, data, semantic):
        table = self.document.add_table(rows=1, cols=len(data["headers"]))
        table.autofit = False
        weights = table_weights(data)
        for column, weight in zip(table.columns, weights):
            column.width = Pt(self.theme["width"] * weight)
        rows = [data["headers"], *data["rows"]]
        compact = len(rows) <= 5 and sum(len(plain(cell)) for row in rows for cell in row) <= 600
        for index, row in enumerate(rows):
            cells = table.rows[0].cells if index == 0 else table.add_row().cells
            properties = table.rows[index]._tr.get_or_add_trPr()
            if index == 0:
                properties.append(element("tblHeader"))
            if sum(len(plain(v)) for v in row) < 1200:
                properties.append(element("cantSplit"))
            for cell, value, weight in zip(cells, row, weights):
                cell.width = Pt(self.theme["width"] * weight)
                p = cell.paragraphs[0]
                p.paragraph_format.line_spacing = 1.25
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.keep_with_next = compact and index < len(rows) - 1
                self.rich(p, value)
                for run in p.runs:
                    run.font.size = Pt(self.theme["table_size"])
                    if index == 0:
                        run.bold = True
                props = cell._tc.get_or_add_tcPr()
                props.append(element("vAlign", val="center"))
                margins = element("tcMar")
                for side in ("top", "bottom", "left", "right"):
                    margins.append(element(side, w=90, type="dxa"))
                props.append(margins)
                props.append(element("shd", fill=self.theme["table_header"][1:] if index == 0 else "F7F8F8" if index % 2 == 0 else "FFFFFF"))
                borders = element("tcBorders")
                for side in ("top", "bottom", "left", "right"):
                    borders.append(element(side, val="single", sz=4, color="D9D9D9"))
                props.append(borders)
        self.document.add_paragraph().paragraph_format.space_after = Pt(2)
        if data.get("caption"):
            self.paragraph(data["caption"], style="Caption")

    def render(self, events, path):
        self.cover()
        toc = navigation(self.model, events)
        if toc:
            self.document.add_heading("Contents", 1)
            for heading in toc:
                p = self.paragraph([{"text": heading["text"], "href": "#" + heading["id"]}])
                p.paragraph_format.left_indent = Pt((heading["level"] - 1) * 12)
            self.document.add_page_break()
        groups = []
        for event in events:
            kind, data, semantic = event["kind"], event["data"], event["semantic"]
            if kind == "group_start":
                groups.append((len(self.document.element.body) - 1, data["keep"]))
                continue
            if kind == "group_end":
                start, keep = groups.pop()
                items = list(self.document.element.body)[start:-1]
                if keep and len(items) <= 8 and all(item.tag == qn("w:p") for item in items):
                    from docx.text.paragraph import Paragraph
                    paragraphs = [Paragraph(item, self.document._body) for item in items]
                    if sum(len(p.text) for p in paragraphs) <= 900:
                        for p in paragraphs[:-1]:
                            p.paragraph_format.keep_with_next = True
                continue
            if kind == "heading":
                p = self.document.add_heading(data["text"], data["level"])
                self.anchor(p, data["id"])
            elif kind == "anchor":
                p = self.document.add_paragraph()
                p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = Pt(1)
                self.anchor(p, data["id"])
            elif kind == "label":
                p = self.paragraph(data["text"], semantic)
                p.paragraph_format.keep_with_next = True
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(5)
                p.runs[0].bold = True
                if semantic:
                    colors = self.theme["semantics"][semantic]
                    p.runs[0].font.color.rgb = RGBColor.from_string(colors["accent"][1:])
                    shade(p, colors["background"], colors["accent"])
                self.anchor(p, data.get("id"))
            elif kind in ("paragraph", "caption", "quote"):
                p = self.paragraph(data["content"], semantic, "Caption" if kind == "caption" else None)
                if kind == "quote":
                    p.paragraph_format.left_indent = Pt(18)
                    for run in p.runs:
                        run.italic = True
                    if data.get("attribution"):
                        self.paragraph(data["attribution"], style="Caption")
            elif kind == "list":
                for index, item in enumerate(data["items"], 1):
                    marker = f"{chr(64 + index)}. " if data.get("letters") else f"{index}. " if data.get("ordered") else "\u2022 "
                    p = self.paragraph(marker, semantic)
                    self.rich(p, item)
                    p.paragraph_format.left_indent = Pt(18)
                    p.paragraph_format.first_line_indent = Pt(-12)
                    p.paragraph_format.space_after = Pt(4)
            elif kind == "code":
                if data.get("language"):
                    self.paragraph(data["language"], style="Caption")
                p = self.paragraph(data["code"], semantic)
                p.paragraph_format.line_spacing = 1.15
                for run in p.runs:
                    run.font.name = "DejaVu Sans Mono"
                    run.font.size = Pt(self.theme["code_size"])
                # Discretionary breaks leave all source characters present and indentation intact.
                for run in p.runs:
                    run.text = "\n".join("\u200b".join(line[i:i + 32] for i in range(0, len(line), 32)) for line in run.text.expandtabs(4).split("\n"))
                shade(p, self.theme["code_background"])
                if data.get("caption"):
                    self.paragraph(data["caption"], style="Caption")
            elif kind == "equation":
                lines = data.get("lines", [{"latex": data.get("latex")}])
                for index, line in enumerate(lines):
                    equation = self.equations.get(line["latex"])
                    width, height, _ = fitted(equation, self.theme["width"] - 55)
                    p = self.document.add_paragraph()
                    p.paragraph_format.keep_with_next = True
                    p.add_run().add_picture(io.BytesIO(equation["png"]), width=Pt(width), height=Pt(height))
                    docpr = p._p.xpath(".//wp:docPr")[0]
                    docpr.set("descr", data["description"])
                    if index == len(lines) - 1 and data.get("number"):
                        p.add_run("    (" + data["number"] + ")")
                    if line.get("explanation"):
                        self.paragraph(line["explanation"], semantic)
                self.paragraph(data["description"], semantic)
                for symbol in data.get("symbols", []):
                    self.paragraph(symbol["symbol"] + ": " + symbol["meaning"], semantic)
            elif kind == "table":
                self.table(data, semantic)
            elif kind == "image":
                filename = image_path(data, self.base)
                with Image.open(filename) as image:
                    ratio = image.height / image.width
                width = min(self.theme["width"] * data.get("width_fraction", 1), 320 / ratio)
                p = self.document.add_paragraph()
                p.paragraph_format.keep_with_next = bool(data.get("caption") or data.get("explanation"))
                p.add_run().add_picture(str(filename), width=Pt(width))
                p._p.xpath(".//wp:docPr")[0].set("descr", data["alt"])
                for key in ("caption", "source_reference", "explanation"):
                    if data.get(key):
                        self.paragraph(data[key], semantic, "Caption" if key != "explanation" else None)
            elif kind == "page_break":
                self.document.add_page_break()
            elif kind == "attempt_space":
                p = self.document.add_paragraph()
                p.paragraph_format.space_after = Pt(12)
            else:
                raise ValueError(f"Unsupported Word event: {kind}")
        self.document.save(path)
